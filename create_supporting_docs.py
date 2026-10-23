"""
Generator for:
1. MedSkill_AI_User_Testing_Form.docx
2. MedSkill_AI_Student_Feedback_Survey.docx
3. MedSkill_AI_Final_Submission_Checklist.docx
4. MedSkill_AI_Technical_Demo_Guide.docx
5. MedSkill_AI_Presentation_Script.docx
"""

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def apply_table_styles(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for j, cell in enumerate(row.cells):
            tcPr = cell._element.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            if i == 0:
                set_cell_background(cell, "0F172A")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(10)
                        run.font.name = "Times New Roman"
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.font.name = "Times New Roman"
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]

def set_doc_defaults(doc):
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(15, 23, 42)

# =====================================================================
# 1. USER TESTING FORM
# =====================================================================
def create_user_testing_form():
    doc = Document()
    set_doc_defaults(doc)
    
    title = doc.add_paragraph()
    r = title.add_run("MedSkill AI — Practical Viva Trainer\nUSER OBSERVATION & TESTING PROTOCOL")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.name = "Times New Roman"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Department of Robotics & Automation | Academic Year 2026-27 | Batch 2024-28 (Semester V)\nHuman-Centered Evaluation Protocol for Clinical Usability")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(18)

    doc.add_heading("1. Participant Demographics & Consent", level=1)
    p = doc.add_paragraph("This observation form is strictly for evaluating digital interface usability and conceptual clarity of practical viva revision. No patient health data or personal identifiers are collected. Data is anonymized using Participant Codes (P01–P14).")
    
    demo_table = doc.add_table(rows=5, cols=2)
    demo_data = [
        ("Participant Code:", "P____ (e.g., P01)"),
        ("Academic Program / Specialization:", "[  ] Respiratory Therapy   [  ] Anaesthesia Tech   [  ] Dialysis Tech   [  ] Other Allied Health"),
        ("Current Year of Study:", "[  ] 1st Year     [  ] 2nd Year     [  ] 3rd Year     [  ] Final Year / Intern"),
        ("Testing Device Used:", "[  ] Android Smartphone   [  ] iPhone   [  ] Laptop (Windows/Mac)   [  ] Tablet"),
        ("Testing Environment:", "[  ] Clinical Skills Lab   [  ] Hospital Break Room   [  ] College Library   [  ] Remote / Hostel")
    ]
    for row_idx, (k, v) in enumerate(demo_data):
        demo_table.rows[row_idx].cells[0].paragraphs[0].add_run(k).font.bold = True
        demo_table.rows[row_idx].cells[1].paragraphs[0].add_run(v)
    apply_table_styles(demo_table, [Inches(2.5), Inches(4.0)])
    demo_table.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)
    set_cell_background(demo_table.rows[0].cells[0], "F1F5F9")
    set_cell_background(demo_table.rows[0].cells[1], "FFFFFF")

    doc.add_heading("2. Standardized Testing Tasks & Observation Rubric", level=1)
    doc.add_paragraph("Instructions to Observer: Do not guide the participant unless they are completely blocked for > 60 seconds. Observe natural hesitation, click patterns, and verbal reactions.")

    tasks_table = doc.add_table(rows=8, cols=4)
    tasks_headers = ["Task #", "Task Description", "Success Criteria", "Observer Notes & Friction Identified"]
    for j, h in enumerate(tasks_headers):
        tasks_table.rows[0].cells[j].paragraphs[0].text = h
    
    tasks_rows = [
        ("Task 1", "Open medskill-ai.streamlit.app and state the application's purpose.", "User identifies viva training purpose within 10 seconds without prompting.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nTime: ____ sec"),
        ("Task 2", "Select your specialization and locate a specific curriculum topic.", "Selects specialty card and uses topic filter to narrow down to 12 questions.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nTime: ____ sec"),
        ("Task 3", "Attempt 3 MCQ questions and verify feedback.", "Selects option, clicks 'Check Answer', reads clinical explanation and key points.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nObservations:"),
        ("Task 4", "Navigate from question 1 to question 2 using the action button.", "Locates the primary blue 'Next Question ➡️' button below the feedback block.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nObservations:"),
        ("Task 5", "Switch to 'Quick Viva Mode' and practice one oral response.", "Types concise answer in text area, clicks 'Compare', reviews model response.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nObservations:"),
        ("Task 6", "Finish session and inspect the Results & Performance breakdown.", "Navigates to Results page, locates accuracy %, and reviews topic cards.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nObservations:"),
        ("Task 7", "Identify weak areas and initiate the '⚡ Quick Retry Quiz'.", "Finds the study recommendation section and launches the 1-click retry on missed items.", "[  ] Completed Unassisted\n[  ] Completed with Hint\n[  ] Failed / Confused\nObservations:")
    ]
    for row_idx, r_data in enumerate(tasks_rows, start=1):
        for col_idx, val in enumerate(r_data):
            tasks_table.rows[row_idx].cells[col_idx].paragraphs[0].text = val
    apply_table_styles(tasks_table, [Inches(0.8), Inches(2.2), Inches(2.0), Inches(2.0)])

    doc.add_heading("3. Post-Test Debrief & Qualitative Inquiries", level=1)
    q1 = doc.add_paragraph()
    q1.add_run("Q1. What specific aspect of the interface caused hesitation or confusion?\n").font.bold = True
    q1.add_run("_________________________________________________________________________________\n_________________________________________________________________________________")
    
    q2 = doc.add_paragraph()
    q2.add_run("Q2. Did the 'Key Viva Revision Points' provide the exact terminology examiners require?\n").font.bold = True
    q2.add_run("_________________________________________________________________________________\n_________________________________________________________________________________")

    q3 = doc.add_paragraph()
    q3.add_run("Q3. What single feature would you add or modify before your next practical examination?\n").font.bold = True
    q3.add_run("_________________________________________________________________________________\n_________________________________________________________________________________")

    doc.save("MedSkill_AI_User_Testing_Form.docx")
    print("Saved MedSkill_AI_User_Testing_Form.docx")

# =====================================================================
# 2. STUDENT FEEDBACK SURVEY
# =====================================================================
def create_student_feedback_survey():
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    r = title.add_run("MedSkill AI — Student Evaluation & Usability Survey")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.name = "Times New Roman"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Human-Centered Innovation Challenge (Activity 13) | Department of Robotics & Automation\nAnonymous Usability & Educational Experience Feedback Form")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)

    doc.add_paragraph("Dear Student Colleague,\nThank you for evaluating MedSkill AI (https://medskill-ai.streamlit.app/). This survey captures your genuine user experience to guide iterative engineering improvements. Please provide candid ratings and feedback.")

    doc.add_heading("Part 1: Quantitative Usability Dimensions (Likert Scale 1–5)", level=1)
    doc.add_paragraph("Rating Scale: 1 = Very Poor / Strongly Disagree | 2 = Poor / Disagree | 3 = Neutral / Acceptable | 4 = Good / Agree | 5 = Excellent / Strongly Agree | N/A = Not Used")

    survey_table = doc.add_table(rows=8, cols=7)
    headers = ["Evaluation Criteria", "1", "2", "3", "4", "5", "N/A"]
    for j, h in enumerate(headers):
        survey_table.rows[0].cells[j].paragraphs[0].text = h

    survey_questions = [
        "1. Ease of Navigation: The dashboard and sidebar made it easy to locate my specialty and topics.",
        "2. Question Clarity: Viva questions were technically accurate and relevant to my university curriculum.",
        "3. Quality of Rationales: The clinical explanations clearly explained the underlying physical laws and mechanisms.",
        "4. Key Viva Bullet Points: The revision expander gave concise buzzwords useful for viva recall.",
        "5. Quick Viva Self-Evaluation: Comparing my typed answer with the model response helped me structure answers.",
        "6. Performance & Retry Flow: The results page accurately identified my weak topics and made retrying simple.",
        "7. Overall Usability: The application was responsive, legible, and easy to use without instruction."
    ]
    for row_idx, q_text in enumerate(survey_questions, start=1):
        survey_table.rows[row_idx].cells[0].paragraphs[0].text = q_text
        for c in range(1, 7):
            survey_table.rows[row_idx].cells[c].paragraphs[0].text = "[  ]"
            survey_table.rows[row_idx].cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styles(survey_table, [Inches(3.5), Inches(0.5), Inches(0.5), Inches(0.5), Inches(0.5), Inches(0.5), Inches(0.6)])

    doc.add_heading("Part 2: Qualitative Critical Feedback", level=1)
    
    p1 = doc.add_paragraph()
    p1.add_run("1. Most Confusing or Frustrating Feature:\n").font.bold = True
    p1.add_run("Which specific interaction felt unnatural, unreadable, or took too long to figure out?\n")
    p1.add_run("_________________________________________________________________________________\n_________________________________________________________________________________")

    p2 = doc.add_paragraph()
    p2.add_run("2. Most High-Yield or Useful Feature:\n").font.bold = True
    p2.add_run("Which tool provided the highest immediate value for your viva examination revision?\n")
    p2.add_run("_________________________________________________________________________________\n_________________________________________________________________________________")

    p3 = doc.add_paragraph()
    p3.add_run("3. First Feature Modification You Would Mandate:\n").font.bold = True
    p3.add_run("If our engineering team can make only one major change before final viva exams, what should it be?\n")
    p3.add_run("_________________________________________________________________________________\n_________________________________________________________________________________")

    p4 = doc.add_paragraph()
    p4.add_run("4. Likelihood of Continued Use:\n").font.bold = True
    p4.add_run("Would you use this prototype for your upcoming university practical examination? [  ] Yes   [  ] No   [  ] Maybe\nWhy or why not? ____________________________________________________________________")

    doc.save("MedSkill_AI_Student_Feedback_Survey.docx")
    print("Saved MedSkill_AI_Student_Feedback_Survey.docx")

# =====================================================================
# 3. FINAL SUBMISSION CHECKLIST
# =====================================================================
def create_final_submission_checklist():
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    r = title.add_run("MedSkill AI — University CA2 Evaluation Checklist")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.name = "Times New Roman"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Requirement-to-Evidence Compliance Matrix | Activity 13: Design for the Unseen\nDepartment of Robotics & Automation, Batch 2024-28 (Semester V)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)

    table = doc.add_table(rows=15, cols=4)
    headers = ["Evaluation Criteria", "Official Requirement", "Verification Evidence in MedSkill AI", "Compliance Status"]
    for j, h in enumerate(headers):
        table.rows[0].cells[j].paragraphs[0].text = h

    checklist_items = [
        ("1. Discovery of Unseen Problem (5 Marks)", "Identify unseen, overlooked real problems; minimum 3 problems evaluated.", "Documented 25 real clinical problems from 8 named stakeholders; evaluated 3 distinct institutional challenges; selected viva paralysis.", "[ X ] Fully Compliant"),
        ("2. Real Stakeholder Focus (5 Marks)", "Do not design for team members; real external user evidence required.", "Focused strictly on Allied Health students (Respiratory, Anaesthesia, Dialysis); shadowed ICU/OT posting shifts.", "[ X ] Fully Compliant"),
        ("3. Empathy Depth (5 Marks)", "Minimum 5 meaningful interactions; user journey and empathy mapping.", "16 total interactions (14 students, 2 clinical faculty); includes Persona, Empathy Map (Says/Thinks/Does/Feels), and Journey Map.", "[ X ] Fully Compliant"),
        ("4. Challenge First Idea (8 Marks)", "'Kill your first idea' documented; minimum 10 alternative ideas across 2 techniques.", "Documented failure of first idea (3D VR knob simulator); generated 10 solutions using Brainstorming & SCAMPER decision matrix.", "[ X ] Fully Compliant"),
        ("5. Prototype 1 Development (5 Marks)", "Build low-fidelity fast prototype using available tools.", "Engineered initial Streamlit application with 24 verified questions; established basic MCQ testing and session tracking.", "[ X ] Fully Compliant"),
        ("6. Real User Testing (5 Marks)", "Put Prototype 1 in front of minimum 10 real users; record likes, dislikes, confusion.", "Tested with 14 students (P01–P14); recorded comprehensive feedback matrix highlighting dark-mode text bug and content requests.", "[ X ] Fully Compliant"),
        ("7. Prototype 2 Iteration (2 Marks)", "Evidence of iteration: Prototype 1 -> User Feedback -> Changes -> Prototype 2.", "Documented 8 concrete iterations: Quick Viva mode, 144 questions, 1-click retry quiz, light theme lock, offline fallback engine.", "[ X ] Fully Compliant"),
        ("8. Final Solution Rigor (5 Marks)", "Clear value proposition: who it is for, how it works, why it is better.", "Streamlit web platform deployed live at medskill-ai.streamlit.app; 144 verified curriculum viva questions across 12 topics.", "[ X ] Fully Compliant"),
        ("9. Presentation Storyline (5 Marks)", "5-minute presentation following: We thought -> observed -> discovered -> built -> users said -> final.", "Structured 8-slide presentation deck (MedSkill_AI_Final_Presentation.pptx) and timed speaking script matching exact prompt.", "[ X ] Fully Compliant"),
        ("10. Question Volume Audit", "At least 12 questions per topic across 12 topics (144 total).", "Automated test test_medskill.py asserts len(QUESTIONS_DATA) == 144 with exactly 12 in every topic; 100% verified.", "[ X ] Fully Compliant"),
        ("11. Zero-Key Architecture", "App must function fully without requiring external API keys.", "Local semantic keyword retrieval engine provides verified textbook explanations even if GEMINI_API_KEY is absent.", "[ X ] Fully Compliant"),
        ("12. High-Contrast Mobile Fix", "Resolve white font / invisible text defect on dark-mode devices.", "Configured .streamlit/config.toml light base theme and enforced !important dark navy text on all radio options.", "[ X ] Fully Compliant"),
        ("13. Live Web Deployment", "Deliverable must be working and accessible to evaluators.", "Live at https://medskill-ai.streamlit.app/; source repository maintained at github.com/Tarun-Teja-74/MEDskill-ai.", "[ X ] Fully Compliant"),
        ("14. Ethical Safety Guardrails", "Not a clinical decision support system; no PHI collected.", "Prominent academic prototype disclaimers across app, documentation, and user forms; zero patient data stored.", "[ X ] Fully Compliant")
    ]

    for row_idx, r_data in enumerate(checklist_items, start=1):
        for col_idx, val in enumerate(r_data):
            table.rows[row_idx].cells[col_idx].paragraphs[0].text = val
    apply_table_styles(table, [Inches(1.8), Inches(1.8), Inches(2.2), Inches(1.2)])

    doc.save("MedSkill_AI_Final_Submission_Checklist.docx")
    print("Saved MedSkill_AI_Final_Submission_Checklist.docx")

# =====================================================================
# 4. TECHNICAL DEMO GUIDE
# =====================================================================
def create_technical_demo_guide():
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    r = title.add_run("MedSkill AI — Technical Architecture & Live Demo Guide")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.name = "Times New Roman"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("System Execution, Architecture Overview & Evaluator Demonstration Script")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)

    doc.add_heading("1. Technical Architecture & Component Model", level=1)
    doc.add_paragraph("MedSkill AI is engineered as a zero-database, stateless, lightweight digital learning platform optimized for instant deployment and zero client friction.")

    arch_table = doc.add_table(rows=6, cols=3)
    arch_headers = ["Layer", "Technology / Framework", "Design Role & Engineering Justification"]
    for j, h in enumerate(arch_headers):
        arch_table.rows[0].cells[j].paragraphs[0].text = h
    
    arch_data = [
        ("Application Layer", "Python 3.13 + Streamlit 1.66", "Single-page reactive web framework; eliminates complex frontend/backend synchronization; enables rapid prototype iteration."),
        ("Data & Question Bank", "Structured In-Memory Dictionaries (questions.py)", "144 verified clinical questions divided into 12 topics. Loaded at startup; zero database round-trip latency; deterministic indexing."),
        ("Intelligence Layer", "Dual-Tier Hybrid (Offline-First + Optional LLM)", "Tier 1: Deterministic keyword matching across 144 curriculum items. Tier 2: Google Gemini 1.5 Flash API with prompt guardrails when key configured."),
        ("State Management", "Streamlit Session State Engine", "Tracks active question list, submitted status, scored dictionary, and prevents duplicate score counting upon multiple clicks."),
        ("Deployment & UI", "Streamlit Community Cloud + Custom CSS", "Hosted on public cloud with continuous git integration; custom CSS enforces high-contrast light theme overrides across dark-mode OS.")
    ]
    for row_idx, r_data in enumerate(arch_data, start=1):
        for col_idx, val in enumerate(r_data):
            arch_table.rows[row_idx].cells[col_idx].paragraphs[0].text = val
    apply_table_styles(arch_table, [Inches(1.5), Inches(2.0), Inches(3.5)])

    doc.add_heading("2. Local Setup & Execution Guide", level=1)
    doc.add_paragraph("To run the application locally on any Windows, macOS, or Linux terminal:")
    code_p = doc.add_paragraph()
    code_p.paragraph_format.left_indent = Inches(0.5)
    code_p.add_run(
        "# Step 1: Clone the repository\n"
        "git clone https://github.com/Tarun-Teja-74/MEDskill-ai.git\n"
        "cd MEDskill-ai\n\n"
        "# Step 2: Install minimal dependencies\n"
        "pip install -r requirements.txt\n\n"
        "# Step 3: Run automated test suite\n"
        "python test_medskill.py\n\n"
        "# Step 4: Launch Streamlit server\n"
        "streamlit run app.py\n"
    ).font.name = "Courier New"

    doc.add_heading("3. Step-by-Step Live Evaluator Demonstration Script (3 Minutes)", level=1)
    
    d1 = doc.add_paragraph()
    d1.add_run("Step 1: The Landing Experience (30 Seconds)\n").font.bold = True
    d1.add_run("• Open https://medskill-ai.streamlit.app/ on browser.\n• Point to the live counters: 144 Available Questions across 3 Specialties (Respiratory, Anaesthesia, Dialysis).\n• Highlight zero login barrier: students jump straight into learning during 10-minute clinical tea breaks.")

    d2 = doc.add_paragraph()
    d2.add_run("Step 2: Specialty & Topic Drill-Down (45 Seconds)\n").font.bold = True
    d2.add_run("• Click 'Select Respiratory Therapy'.\n• In topic dropdown, select 'Ventilator fundamentals'. Point out the summary card updating to exactly 12 questions.\n• Select 'MCQ Practice' mode and click 'Begin Viva Practice'.")

    d3 = doc.add_paragraph()
    d3.add_run("Step 3: Answering, Feedback & Buzzwords (60 Seconds)\n").font.bold = True
    d3.add_run("• Select an answer (e.g. Question on Peak vs Plateau pressure).\n• Click 'Check Answer'. Point out the high-contrast card styling and instant clinical explanation.\n• Expand the '📌 Key Viva Revision Points' card: explain that this provides the exact keywords external examiners grade.")

    d4 = doc.add_paragraph()
    d4.add_run("Step 4: Oral Articulation & 1-Click Retry (45 Seconds)\n").font.bold = True
    d4.add_run("• Demonstrate 'Quick Viva Mode': show how typing an oral response forces active verbal reasoning before revealing the model answer.\n• Navigate to 'Results & Feedback': show the topic-wise accuracy cards.\n• Click '⚡ Quick Retry Quiz' to show the system isolating only missed questions for remediation.")

    doc.save("MedSkill_AI_Technical_Demo_Guide.docx")
    print("Saved MedSkill_AI_Technical_Demo_Guide.docx")

# =====================================================================
# 5. PRESENTATION SCRIPT
# =====================================================================
def create_presentation_script():
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    r = title.add_run("MedSkill AI — 5-Minute Final Viva Presentation Script")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.name = "Times New Roman"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Department of Robotics & Automation | CA2 Activity 13: Design for the Unseen\nStrict 5-Minute Timed Delivery Script for Team Presentation")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)

    doc.add_paragraph("Presentation Guidelines: Each team receives exactly 5 minutes. The story must follow the required narrative: We thought -> We observed -> We discovered -> We realised the real problem was -> We generated -> We built -> Users told us -> So we changed -> Our final solution is.")

    script_sections = [
        ("Slide 1: Title & Introduction (0:00 – 0:30)", 
         "Good morning, respected faculty evaluators. Our project for the 'Design for the Unseen' innovation challenge is titled MedSkill AI—an AI-powered practical viva trainer for healthcare technology students. Our team set out to investigate an unseen user group in medical hospitals: undergraduate students in Respiratory Therapy, Anaesthesia Technology, and Dialysis Technology."),
        
        ("Slide 2: We Thought... & Initial Assumptions (0:30 – 1:00)", 
         "When we began, our initial assumption was simple: we thought allied health students were failing or struggling in university practical vivas because they lacked notes, question banks, or revision material. We assumed the solution was building another PDF library."),
        
        ("Slide 3: We Observed & Discovered... (1:00 – 1:45)", 
         "We challenged our assumptions. We went to clinical skills labs and hospital postings, shadowing 8 student technologists: Riddhi, Siddhi, Naziya, Preethi, Bhargav, Naveen, Tharun, and Atharv. We observed 25 authentic clinical bottlenecks—from dry oxygen humidifier bottles and severe alarm fatigue to ventilator circuit disconnects. When we questioned students, we discovered they had hundreds of pages of PDF notes. But during mock vivas, when asked 'Why did peak pressure spike while plateau remained normal?', they froze completely. We discovered the real culprit: Passive Revision Syndrome. Reading textbooks gives an illusion of competence, but does not train verbal recall under examiner pressure."),
        
        ("Slide 4: We Realised the Real Problem Was... (1:45 – 2:15)", 
         "We reframed the problem. The real problem is not a lack of content—it is a lack of an interactive, low-stakes digital simulator that diagnoses conceptual misunderstandings and helps students practice verbal articulation before standing in front of an external examiner."),
        
        ("Slide 5: We Generated & Killed Our First Idea (2:15 – 2:50)", 
         "We used Brainstorming and SCAMPER to generate 10 alternative solutions. Importantly, we killed our first obvious idea: a 3D virtual reality knob-turning simulator. Why? Because students already touch real machines in hospital postings; their real bottleneck is oral physics articulation. Furthermore, VR headsets are expensive and inaccessible. Instead, we prioritized a zero-friction, dual-mode web simulator that runs on any smartphone."),
        
        ("Slide 6: We Built Prototype 1 (2:50 – 3:25)", 
         "We built Prototype 1 using Python and Streamlit, featuring an initial set of 24 multiple-choice questions across the three specialties."),
        
        ("Slide 7: Users Told Us... & Evidence of Iteration (3:25 – 4:15)", 
         "We put Prototype 1 in front of 14 students. They gave us candid, critical feedback: First, Riddhi pointed out that in dark mode on mobile, the radio text was white and completely invisible. Second, Siddhi noted that 2 questions per topic was insufficient. Third, Naziya stated: 'Vivas are oral, not multiple-choice!' Fourth, Naveen pointed out that hospital basements have zero mobile reception. Based on this genuine feedback, we iterated to Prototype 2: we locked in high-contrast light styling, expanded the question bank to 144 verified questions (12 in every single topic), engineered 'Quick Viva Mode' for oral response comparison, added a 1-click retry on missed items, and built an offline-first knowledge engine that requires zero API keys."),
        
        ("Slide 8: Our Final Solution & Impact (4:15 – 5:00)", 
         "Our final solution is live today at medskill-ai.streamlit.app. It provides 144 verified curriculum questions, dual-mode practice, targeted revision recommendations, and complete offline resilience. It is an educational prototype that bridges the gap between hospital floor chaos and viva examination mastery. Thank you!")
    ]

    for s_title, s_text in script_sections:
        h = doc.add_paragraph()
        r = h.add_run(s_title)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(15, 23, 42)
        
        p = doc.add_paragraph(s_text)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(12)

    doc.save("MedSkill_AI_Presentation_Script.docx")
    print("Saved MedSkill_AI_Presentation_Script.docx")

if __name__ == "__main__":
    create_user_testing_form()
    create_student_feedback_survey()
    create_final_submission_checklist()
    create_technical_demo_guide()
    create_presentation_script()
